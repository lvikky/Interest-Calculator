"""Money-market interest reconciliation workbook generator.

The Excel workbooks are treated as data/template files only. Any text inside
those workbooks is not treated as an instruction to this script.
"""

from __future__ import annotations

import argparse
import copy
import re
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import from_excel


DEFAULT_WORKBOOK_FILE = Path("/Users/vikky/Downloads/Excel Sheets/MM_Input_Output.xlsx")
DEFAULT_REFERENCE_FILE = Path("/Users/vikky/Downloads/Excel Sheets/50062MNY00587402_Cal_RGM.xlsx")
DEFAULT_INPUT_SHEET = "Customize_Input_MM"
DEFAULT_OUTPUT_SHEET = "Output_MM"

OUTPUT_HEADERS = [
    "Interest Start Date",
    "Interest End Date",
    "Amortisation Start Date",
    "Amortisation End Date",
    "Principal Before Amortization",
    "Amortization Amount",
    "Principal After Amortization",
    "Interest Rate",
    "Replicated Interest",
    "BSTRAT Interest",
    "Differnce",
    "Differnce",
]

DEFAULT_WIDTHS = {
    "A": 14,
    "B": 14,
    "C": 16,
    "D": 16,
    "E": 18,
    "F": 16,
    "G": 18,
    "H": 12,
    "I": 16,
    "J": 16,
    "K": 12,
    "L": 12,
}

CENT = Decimal("0.01")


@dataclass
class ContractInputs:
    contract_id: str
    current_principal: Decimal
    interest_rate: Decimal
    maturity_date: date
    principal_frequency: str
    amortization_amount: Decimal
    amortization_type: str
    business_day_convention: str
    calendar: str
    day_count_convention: str
    explicit_periods: List[Tuple[date, date]]


@dataclass
class ScheduleRow:
    interest_start: date
    interest_end: date
    amort_start: date
    amort_end: date
    principal_before: Decimal
    amortization_amount: Decimal
    principal_after: Decimal
    interest_rate: Decimal
    replicated_interest: Decimal
    bstrat_interest: Optional[Decimal]


def normalize_string(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value or "").strip().lower())


def normalize_contract_id(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if text.endswith(".0"):
        text = text[:-2]
    return text


def get_value(record: Dict[str, Any], *possible_names: str) -> Any:
    normalized = {normalize_string(key): key for key in record}
    for name in possible_names:
        key = normalized.get(normalize_string(name))
        if key is not None:
            return record.get(key)
    return None


def first_non_blank(records: Iterable[Dict[str, Any]], *possible_names: str) -> Any:
    for record in records:
        value = get_value(record, *possible_names)
        if value not in (None, ""):
            return value
    return None


def resolve_sheet_name(sheet_names: Iterable[str], requested_name: str, label: str) -> str:
    normalized_sheets = {normalize_string(sheet): sheet for sheet in sheet_names}
    match = normalized_sheets.get(normalize_string(requested_name))
    if match is not None:
        return match

    available = ", ".join(sheet_names)
    raise ValueError(
        f"Could not find the {label} worksheet {requested_name!r}. "
        f"Available worksheets: {available or 'none'}"
    )


def worksheet_to_records(worksheet: Any) -> List[Dict[str, Any]]:
    headers = [cell.value for cell in next(worksheet.iter_rows(min_row=1, max_row=1))]
    records: List[Dict[str, Any]] = []

    for row_values in worksheet.iter_rows(min_row=2, values_only=True):
        if all(value in (None, "") for value in row_values):
            continue

        record: Dict[str, Any] = {}
        for idx, header in enumerate(headers):
            if header in (None, ""):
                continue
            record[str(header).strip()] = row_values[idx] if idx < len(row_values) else None
        records.append(record)

    return records


def read_excel_records(file_path: Path, sheet_name: Optional[str] = None) -> List[Dict[str, Any]]:
    if not file_path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    workbook = load_workbook(file_path, data_only=True, read_only=True)
    worksheet = workbook[sheet_name] if sheet_name else workbook.worksheets[0]
    records = worksheet_to_records(worksheet)
    workbook.close()
    return records


def read_source_workbook(
    workbook_file: Path,
    input_sheet: str,
    output_sheet: str,
) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]]]:
    if not workbook_file.exists():
        raise FileNotFoundError(f"Workbook not found: {workbook_file}")

    workbook = load_workbook(workbook_file, data_only=True, read_only=True)
    try:
        input_sheet_name = resolve_sheet_name(workbook.sheetnames, input_sheet, "input")
        output_sheet_name = resolve_sheet_name(workbook.sheetnames, output_sheet, "output")

        input_records = worksheet_to_records(workbook[input_sheet_name])
        output_records = worksheet_to_records(workbook[output_sheet_name])
    finally:
        workbook.close()

    return input_records, output_records


def parse_decimal(value: Any, field_name: str) -> Decimal:
    if value in (None, ""):
        raise ValueError(f"Missing required numeric value: {field_name}")

    text = str(value).strip().replace(",", "")
    if text.startswith("(") and text.endswith(")"):
        text = f"-{text[1:-1]}"

    try:
        return Decimal(text)
    except InvalidOperation as exc:
        raise ValueError(f"Invalid numeric value for {field_name}: {value!r}") from exc


def parse_date(value: Any, field_name: str) -> date:
    if value in (None, ""):
        raise ValueError(f"Missing required date value: {field_name}")

    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        return from_excel(value).date()

    text = str(value).strip()
    for fmt in ("%m/%d/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass

    raise ValueError(f"Invalid date value for {field_name}: {value!r}")


def excel_date_text(value: date) -> str:
    return f"{value.month}/{value.day}/{value.year}"


def round_money(value: Decimal) -> Decimal:
    return value.quantize(CENT, rounding=ROUND_HALF_UP)


def filter_contract_rows(records: List[Dict[str, Any]], contract_id: str, id_field: str) -> List[Dict[str, Any]]:
    wanted = normalize_contract_id(contract_id)
    return [
        record
        for record in records
        if normalize_contract_id(get_value(record, id_field)) == wanted
    ]


def parse_frequency_months(frequency: Any) -> int:
    if frequency in (None, ""):
        raise ValueError("Missing PRINCIPAL_FREQUENCY. Cannot build amortization periods.")

    text = str(frequency).strip().upper()
    match = re.fullmatch(r"(\d+)\s*M", text)
    if match:
        return int(match.group(1))

    if text in {"M", "MONTH", "MONTHLY"}:
        return 1
    if text in {"Q", "QTR", "QUARTER", "QUARTERLY"}:
        return 3

    if text.isdigit():
        months = int(text)
        if months > 0:
            return months

    raise ValueError(f"Unsupported PRINCIPAL_FREQUENCY {frequency!r}. This script supports monthly frequencies such as 1M.")


def add_months(value: date, months: int) -> date:
    month_index = value.month - 1 + months
    year = value.year + month_index // 12
    month = month_index % 12 + 1
    day = min(value.day, last_day_of_month(year, month))
    return date(year, month, day)


def last_day_of_month(year: int, month: int) -> int:
    if month == 12:
        next_month = date(year + 1, 1, 1)
    else:
        next_month = date(year, month + 1, 1)
    return (next_month - timedelta(days=1)).day


def nth_weekday(year: int, month: int, weekday: int, n: int) -> date:
    current = date(year, month, 1)
    while current.weekday() != weekday:
        current += timedelta(days=1)
    return current + timedelta(days=7 * (n - 1))


def last_weekday(year: int, month: int, weekday: int) -> date:
    current = date(year, month, last_day_of_month(year, month))
    while current.weekday() != weekday:
        current -= timedelta(days=1)
    return current


def observed_fixed_holiday(year: int, month: int, day: int) -> date:
    holiday = date(year, month, day)
    if holiday.weekday() == 5:
        return holiday - timedelta(days=1)
    if holiday.weekday() == 6:
        return holiday + timedelta(days=1)
    return holiday


def us_bank_holidays(year: int) -> set[date]:
    return {
        observed_fixed_holiday(year, 1, 1),
        nth_weekday(year, 1, 0, 3),
        nth_weekday(year, 2, 0, 3),
        last_weekday(year, 5, 0),
        observed_fixed_holiday(year, 6, 19),
        observed_fixed_holiday(year, 7, 4),
        nth_weekday(year, 9, 0, 1),
        nth_weekday(year, 10, 0, 2),
        observed_fixed_holiday(year, 11, 11),
        nth_weekday(year, 11, 3, 4),
        observed_fixed_holiday(year, 12, 25),
    }


def is_business_day(value: date, calendar: str) -> bool:
    if value.weekday() >= 5:
        return False

    calendar_key = normalize_string(calendar)
    if calendar_key in {"nyb", "us", "usd", "usbank"}:
        return value not in us_bank_holidays(value.year)

    return True


def next_business_day(value: date, calendar: str) -> date:
    current = value
    while not is_business_day(current, calendar):
        current += timedelta(days=1)
    return current


def previous_business_day(value: date, calendar: str) -> date:
    current = value
    while not is_business_day(current, calendar):
        current -= timedelta(days=1)
    return current


def adjust_business_day(value: date, convention: str, calendar: str) -> date:
    key = normalize_string(convention)
    if key in {"", "same", "none", "noadjustment"}:
        return value
    if is_business_day(value, calendar):
        return value

    if key in {"next", "following"}:
        return next_business_day(value, calendar)
    if key in {"prev", "previous", "preceding"}:
        return previous_business_day(value, calendar)

    # NEXT_PREV in the input behaves like modified following in the reference workbook:
    # use the next business day unless that crosses into a new month, then use previous.
    if key in {"nextprev", "modifiedfollowing", "modfollowing"}:
        candidate = next_business_day(value, calendar)
        if candidate.month == value.month:
            return candidate
        return previous_business_day(value, calendar)

    raise ValueError(f"Unsupported business day convention: {convention!r}")


def day_count_days(start_date: date, end_date: date, convention: str) -> int:
    key = normalize_string(convention)
    if "30360" in key:
        start_day = min(start_date.day, 30)
        end_day = end_date.day
        if end_day == 31 and start_day == 30:
            end_day = 30
        return (
            360 * (end_date.year - start_date.year)
            + 30 * (end_date.month - start_date.month)
            + (end_day - start_day)
        )

    return (end_date - start_date).days


def day_count_denominator(convention: str) -> Decimal:
    key = normalize_string(convention)
    if "365" in key:
        return Decimal("365")
    return Decimal("360")


def build_contract_inputs(contract_rows: List[Dict[str, Any]], contract_id: str) -> ContractInputs:
    if not contract_rows:
        raise ValueError(f"Contract ID {contract_id!r} was not found in the input workbook.")

    rows = list(contract_rows)
    explicit_periods: List[Tuple[date, date]] = []
    for row in rows:
        start_raw = get_value(row, "PRINCIPAL_AMORT_START_DATE", "Amortisation Start Date")
        end_raw = get_value(row, "PRINCIPAL_AMORT_END_DATE", "Amortisation End Date")
        if start_raw not in (None, "") and end_raw not in (None, ""):
            explicit_periods.append((
                parse_date(start_raw, "PRINCIPAL_AMORT_START_DATE"),
                parse_date(end_raw, "PRINCIPAL_AMORT_END_DATE"),
            ))

    explicit_periods.sort(key=lambda pair: pair[1])
    if not explicit_periods:
        raise ValueError(
            f"Contract ID {contract_id!r} has no PRINCIPAL_AMORT_START_DATE / "
            "PRINCIPAL_AMORT_END_DATE rows in the input workbook."
        )

    return ContractInputs(
        contract_id=normalize_contract_id(contract_id),
        current_principal=parse_decimal(first_non_blank(rows, "CURRENT_PRINCIPAL", "ORIGINAL_PRINCIPAL"), "CURRENT_PRINCIPAL"),
        interest_rate=parse_decimal(first_non_blank(rows, "CURRENT_INTEREST_RATE", "FEE_RATE", "INTEREST_RATE"), "CURRENT_INTEREST_RATE"),
        maturity_date=parse_date(first_non_blank(rows, "MATURITY_DATE"), "MATURITY_DATE"),
        principal_frequency=str(first_non_blank(rows, "PRINCIPAL_FREQUENCY") or "").strip(),
        amortization_amount=parse_decimal(first_non_blank(rows, "PRINCIPAL_AMORT_AMOUNT"), "PRINCIPAL_AMORT_AMOUNT"),
        amortization_type=str(first_non_blank(rows, "PRINCIPAL_AMORTIZATION_TYPE") or "RGM").strip(),
        business_day_convention=str(first_non_blank(rows, "BDC") or "SAME").strip(),
        calendar=str(first_non_blank(rows, "CALENDAR") or "").strip(),
        day_count_convention=str(first_non_blank(rows, "INTEREST_DAY_COUNT_CONVENTION") or "Actual/360").strip(),
        explicit_periods=explicit_periods,
    )


def generate_amortization_periods(inputs: ContractInputs) -> List[Tuple[date, date]]:
    months = parse_frequency_months(inputs.principal_frequency)
    seed_start = inputs.explicit_periods[0][0]
    periods: List[Tuple[date, date]] = []
    period_number = 1

    while True:
        if period_number <= len(inputs.explicit_periods):
            amort_start, amort_end = inputs.explicit_periods[period_number - 1]
        else:
            raw_end = add_months(seed_start, months * period_number)
            if raw_end > inputs.maturity_date:
                amort_end = adjust_business_day(inputs.maturity_date, inputs.business_day_convention, inputs.calendar)
            else:
                amort_end = adjust_business_day(raw_end, inputs.business_day_convention, inputs.calendar)
            amort_start = add_months(amort_end, -months)

        periods.append((amort_start, amort_end))
        if amort_end >= inputs.maturity_date:
            break

        period_number += 1
        if period_number > 1200:
            raise RuntimeError("Generated too many amortization periods. Check the maturity date and frequency.")

    return periods


def build_output_interest_lookup(output_rows: List[Dict[str, Any]], contract_id: str) -> Dict[date, Decimal]:
    interest_by_date: Dict[date, Decimal] = defaultdict(lambda: Decimal("0"))

    for row in output_rows:
        if normalize_contract_id(get_value(row, "UNIQUE_ID", "CONTRACT_ID")) != normalize_contract_id(contract_id):
            continue

        component = str(get_value(row, "COMPONENT") or "").strip().upper()
        if component and component != "IR":
            continue

        payment_date = parse_date(get_value(row, "FLOW_PAYMENT_DATE", "PAYMENT_DATE"), "FLOW_PAYMENT_DATE")
        amount = parse_decimal(get_value(row, "AMOUNT", "AMOUNT_USD", "BSTRAT Interest"), "AMOUNT")
        interest_by_date[payment_date] += amount

    return dict(interest_by_date)


def calculate_replicated_interest(inputs: ContractInputs, principal: Decimal, start_date: date, end_date: date) -> Decimal:
    rate_decimal = inputs.interest_rate / Decimal("100") if inputs.interest_rate >= 1 else inputs.interest_rate
    days = Decimal(day_count_days(start_date, end_date, inputs.day_count_convention))
    denominator = day_count_denominator(inputs.day_count_convention)
    return round_money(-(principal * rate_decimal * days / denominator))


def build_schedule(inputs: ContractInputs, output_interest: Dict[date, Decimal]) -> List[ScheduleRow]:
    periods = generate_amortization_periods(inputs)
    schedule: List[ScheduleRow] = []
    previous_interest_end: Optional[date] = None
    principal_before = round_money(inputs.current_principal)

    for amort_start, amort_end in periods:
        interest_start = previous_interest_end or amort_start
        interest_end = amort_end
        if amort_end >= inputs.maturity_date:
            amortization_amount = principal_before
        else:
            amortization_amount = min(round_money(inputs.amortization_amount), principal_before)
        principal_after = round_money(principal_before - amortization_amount)
        replicated_interest = calculate_replicated_interest(inputs, principal_before, interest_start, interest_end)
        bstrat_interest = output_interest.get(interest_end)

        schedule.append(
            ScheduleRow(
                interest_start=interest_start,
                interest_end=interest_end,
                amort_start=amort_start,
                amort_end=amort_end,
                principal_before=principal_before,
                amortization_amount=amortization_amount,
                principal_after=principal_after,
                interest_rate=inputs.interest_rate,
                replicated_interest=replicated_interest,
                bstrat_interest=round_money(bstrat_interest) if bstrat_interest is not None else None,
            )
        )

        principal_before = principal_after
        previous_interest_end = interest_end
        if principal_before <= Decimal("0"):
            break

    return schedule


def copy_cell_style(source_cell: Any, target_cell: Any) -> None:
    if source_cell.has_style:
        target_cell.font = copy.copy(source_cell.font)
        target_cell.fill = copy.copy(source_cell.fill)
        target_cell.border = copy.copy(source_cell.border)
        target_cell.alignment = copy.copy(source_cell.alignment)
        target_cell.protection = copy.copy(source_cell.protection)
        target_cell.number_format = source_cell.number_format


def apply_fallback_style(worksheet: Any, row_count: int) -> None:
    header_fill = PatternFill("solid", fgColor="D9E1F2")
    thin = Side(style="thin", color="000000")
    medium = Side(style="medium", color="000000")

    for col_idx in range(1, len(OUTPUT_HEADERS) + 1):
        header = worksheet.cell(row=1, column=col_idx)
        header.fill = header_fill
        header.font = Font(name="Calibri", size=11, bold=True, color="000000")
        header.alignment = Alignment(horizontal="center", vertical="center")
        header.border = Border(left=thin, right=thin, top=thin, bottom=medium)

    for row_idx in range(2, row_count + 2):
        for col_idx in range(1, len(OUTPUT_HEADERS) + 1):
            cell = worksheet.cell(row=row_idx, column=col_idx)
            cell.font = Font(name="Calibri", size=11)
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            cell.alignment = Alignment(horizontal="right", vertical="center")

        for col_idx in (1, 2, 3, 4):
            worksheet.cell(row=row_idx, column=col_idx).alignment = Alignment(horizontal="center", vertical="center")

        for col_idx in (5, 6, 7):
            worksheet.cell(row=row_idx, column=col_idx).number_format = "#,##0.00"
        worksheet.cell(row=row_idx, column=8).number_format = "0.000"
        for col_idx in (9, 10, 11, 12):
            worksheet.cell(row=row_idx, column=col_idx).number_format = "#,##0.00;(#,##0.00);0.00"


def apply_reference_style(worksheet: Any, reference_file: Path, row_count: int) -> Optional[List[str]]:
    if not reference_file.exists():
        return None

    reference_workbook = load_workbook(reference_file, data_only=False)
    reference_sheet = reference_workbook.worksheets[0]
    max_col = min(reference_sheet.max_column, len(OUTPUT_HEADERS))

    headers = [reference_sheet.cell(row=1, column=col_idx).value for col_idx in range(1, max_col + 1)]

    for col_idx in range(1, max_col + 1):
        letter = get_column_letter(col_idx)
        width = reference_sheet.column_dimensions[letter].width
        if width:
            worksheet.column_dimensions[letter].width = width
        copy_cell_style(reference_sheet.cell(row=1, column=col_idx), worksheet.cell(row=1, column=col_idx))

    for row_idx in range(2, row_count + 2):
        template_row = min(row_idx, reference_sheet.max_row)
        for col_idx in range(1, max_col + 1):
            copy_cell_style(
                reference_sheet.cell(row=template_row, column=col_idx),
                worksheet.cell(row=row_idx, column=col_idx),
            )

    reference_workbook.close()
    return [str(header) if header is not None else "" for header in headers]


def safe_sheet_name(name: str) -> str:
    cleaned = re.sub(r"[\[\]:*?/\\]", "_", name)
    return cleaned[:31] or "Interest_Calculation"


def unique_output_path(output_dir: Path, file_name: str, overwrite: bool) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    candidate = output_dir / file_name
    if overwrite or not candidate.exists():
        return candidate

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return output_dir / f"{candidate.stem}_{timestamp}{candidate.suffix}"


def decimal_to_float(value: Optional[Decimal]) -> Optional[float]:
    if value is None:
        return None
    return float(value)


def export_schedule(
    inputs: ContractInputs,
    schedule: List[ScheduleRow],
    output_dir: Path,
    reference_file: Path,
    overwrite: bool,
) -> Path:
    sheet_name = safe_sheet_name(f"{inputs.contract_id}_Cal_{inputs.amortization_type or 'RGM'}")
    output_path = unique_output_path(output_dir, f"{sheet_name}.xlsx", overwrite=overwrite)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name

    headers = apply_reference_style(worksheet, reference_file, len(schedule)) or OUTPUT_HEADERS
    for col_idx, header in enumerate(headers, start=1):
        worksheet.cell(row=1, column=col_idx, value=header)

    if headers == OUTPUT_HEADERS:
        for letter, width in DEFAULT_WIDTHS.items():
            worksheet.column_dimensions[letter].width = width
        apply_fallback_style(worksheet, len(schedule))

    for row_idx, item in enumerate(schedule, start=2):
        values = [
            excel_date_text(item.interest_start),
            excel_date_text(item.interest_end),
            excel_date_text(item.amort_start),
            excel_date_text(item.amort_end),
            decimal_to_float(item.principal_before),
            decimal_to_float(item.amortization_amount),
            f"=E{row_idx}-F{row_idx}" if item.principal_after > Decimal("0") else None,
            decimal_to_float(item.interest_rate),
            decimal_to_float(item.replicated_interest),
            decimal_to_float(item.bstrat_interest),
            f"=I{row_idx}-J{row_idx}" if item.bstrat_interest is not None else None,
            None,
        ]
        for col_idx, value in enumerate(values, start=1):
            worksheet.cell(row=row_idx, column=col_idx, value=value)

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.save(output_path)
    return output_path


def run_reconciliation(
    workbook_file: Path,
    input_sheet: str,
    output_sheet: str,
    reference_file: Path,
    output_dir: Path,
    contract_id: str,
    overwrite: bool = False,
) -> Path:
    input_records, output_records = read_source_workbook(
        workbook_file=workbook_file,
        input_sheet=input_sheet,
        output_sheet=output_sheet,
    )

    contract_rows = filter_contract_rows(input_records, contract_id, "CONTRACT_ID")
    inputs = build_contract_inputs(contract_rows, contract_id)

    output_interest = build_output_interest_lookup(output_records, contract_id)
    if not output_interest:
        available_ids = sorted(
            {
                normalize_contract_id(get_value(row, "UNIQUE_ID", "CONTRACT_ID"))
                for row in output_records
                if normalize_contract_id(get_value(row, "UNIQUE_ID", "CONTRACT_ID"))
            }
        )
        raise ValueError(
            f"No IR rows were found in worksheet {output_sheet!r} for Contract ID {contract_id!r}. "
            f"Available output IDs: {', '.join(available_ids) or 'none'}"
        )

    schedule = build_schedule(inputs, output_interest)
    missing_dates = [row.interest_end for row in schedule if row.bstrat_interest is None]
    if missing_dates:
        missing_text = ", ".join(excel_date_text(value) for value in missing_dates[:10])
        extra = "..." if len(missing_dates) > 10 else ""
        print(f"[!] BSTRAT interest missing for {len(missing_dates)} schedule date(s): {missing_text}{extra}")

    return export_schedule(inputs, schedule, output_dir, reference_file, overwrite=overwrite)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Generate a money-market interest reconciliation workbook for one contract."
    )
    parser.add_argument(
        "--workbook-file",
        type=Path,
        default=DEFAULT_WORKBOOK_FILE,
        help="Path to one workbook containing both input and output worksheets.",
    )
    parser.add_argument(
        "--input-sheet",
        default=DEFAULT_INPUT_SHEET,
        help="Worksheet containing input contract rows.",
    )
    parser.add_argument(
        "--output-sheet",
        default=DEFAULT_OUTPUT_SHEET,
        help="Worksheet containing output cash flow rows.",
    )
    parser.add_argument(
        "--reference-file",
        type=Path,
        default=DEFAULT_REFERENCE_FILE,
        help="Reference workbook used only for output headers and styling.",
    )
    parser.add_argument("--output-dir", type=Path, default=Path.cwd(), help="Folder where the generated workbook is saved")
    parser.add_argument("--contract-id", help="Contract ID to process. If omitted, the script prompts for it.")
    parser.add_argument("--overwrite", action="store_true", help="Overwrite the generated workbook if it already exists")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    contract_id = normalize_contract_id(args.contract_id or input("Enter Contract ID: "))
    if not contract_id:
        raise ValueError("Contract ID is required.")

    print(f"[*] Reading source workbook: {args.workbook_file}")
    print(f"[*] Input worksheet: {args.input_sheet}")
    print(f"[*] Output worksheet: {args.output_sheet}")
    print(f"[*] Processing contract ID: {contract_id}")

    output_path = run_reconciliation(
        workbook_file=args.workbook_file,
        input_sheet=args.input_sheet,
        output_sheet=args.output_sheet,
        reference_file=args.reference_file,
        output_dir=args.output_dir,
        contract_id=contract_id,
        overwrite=args.overwrite,
    )

    print(f"[+] Workbook generated: {output_path}")


if __name__ == "__main__":
    main()
