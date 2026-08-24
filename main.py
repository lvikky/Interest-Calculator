import os
import re
from datetime import datetime
from typing import Any, Dict, List, Optional, Tuple, Union

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


# ---------------------------------------------------------
# 1. Helper & Column Normalization Functions
# ---------------------------------------------------------
def normalize_string(val: str) -> str:
    """Normalizes string headers by trimming, lowercasing, and removing non-alphanumeric characters."""
    return re.sub(r'[^a-z0-9]+', '', str(val).strip().lower())

def find_column(df: pd.DataFrame, possible_names: List[str]) -> Optional[str]:
    """Finds a matching column in a DataFrame given a list of candidate names (case and punctuation insensitive)."""
    normalized_cols = {normalize_string(col): col for col in df.columns}
    for candidate in possible_names:
        norm_cand = normalize_string(candidate)
        if norm_cand in normalized_cols:
            return normalized_cols[norm_cand]
    return None


def resolve_sheet_name(sheet_names: List[str], *preferred_names: str) -> Optional[str]:
    """Resolves matching sheet name from list of sheet names."""
    sheet_map = {normalize_string(sheet): sheet for sheet in sheet_names}
    for preferred in preferred_names:
        norm = normalize_string(preferred)
        if norm in sheet_map:
            return sheet_map[norm]
    for sheet in sheet_names:
        norm = normalize_string(sheet)
        if any(token in norm for token in ('contract', 'input', 'cashflow', 'master')):
            return sheet
    return None


def normalize_contract_id(val: Any) -> str:
    """Normalizes contract ID by stripping whitespace and trailing .0 if loaded as float."""
    if pd.isna(val) or val is None:
        return ""
    s = str(val).strip()
    if s.endswith('.0'):
        s = s[:-2]
    return s


def parse_frequency_offset(freq_val: Any, end_date: pd.Timestamp) -> pd.Timestamp:
    """Calculates start date by subtracting INTEREST_FREQUENCY from first cashflow end date."""
    if pd.isna(freq_val) or freq_val is None:
        return end_date - pd.DateOffset(months=3)

    freq_str = str(freq_val).strip().lower()

    # If numeric (e.g., 90, 30, 3)
    try:
        num_val = float(freq_str)
        if num_val > 15:
            return end_date - pd.Timedelta(days=int(num_val))
        elif num_val > 0:
            return end_date - pd.DateOffset(months=int(num_val))
    except ValueError:
        pass

    # String representation checks
    if 'quarter' in freq_str or freq_str in ('q', 'qtr', 'quarterly'):
        return end_date - pd.DateOffset(months=3)
    elif 'month' in freq_str or freq_str in ('m', 'monthly'):
        return end_date - pd.DateOffset(months=1)
    elif 'semi' in freq_str or 'half' in freq_str or freq_str in ('sa', 'semi-annual', 'semi-annually'):
        return end_date - pd.DateOffset(months=6)
    elif 'annual' in freq_str or 'year' in freq_str or freq_str in ('a', 'y', 'annually'):
        return end_date - pd.DateOffset(years=1)

    return end_date - pd.DateOffset(months=3)


def get_frequency_days(freq_val: Any) -> int:
    """Extracts payment frequency in days from input table (defaults to 90 if unspecified/quarterly)."""
    if pd.isna(freq_val) or freq_val is None:
        return 90

    freq_str = str(freq_val).strip().lower()

    # If numeric (e.g., 90, 30, 180, 360, or 3, 1, 6, 12)
    try:
        num_val = float(freq_str)
        if num_val > 15:
            return int(num_val)
        elif num_val == 3:
            return 90
        elif num_val == 1:
            return 30
        elif num_val == 6:
            return 180
        elif num_val == 12:
            return 360
    except ValueError:
        pass

    # String representation checks
    if 'quarter' in freq_str or freq_str in ('q', 'qtr', 'quarterly'):
        return 90
    elif 'month' in freq_str or freq_str in ('m', 'monthly'):
        return 30
    elif 'semi' in freq_str or 'half' in freq_str or freq_str in ('sa', 'semi-annual', 'semi-annually'):
        return 180
    elif 'annual' in freq_str or 'year' in freq_str or freq_str in ('a', 'y', 'annually'):
        return 360

    return 90


# ---------------------------------------------------------
# 2. Financial Ingestion & Mapping Layer
# ---------------------------------------------------------
def load_financial_data(
    input_source: str = "BSTRAT_Cashflow_Analysis.xlsx",
    cashflow_source: Optional[str] = "Output_BSTRAT_Cashflow.xlsx",
    target_contract_id: Optional[str] = None
) -> Tuple[pd.DataFrame, pd.DataFrame]:
    """
    Loads contract master parameters and cashflow event records.
    Supports either two separate files, or a single Excel workbook containing both sheets.
    """
    input_df: Optional[pd.DataFrame] = None
    cashflow_df: Optional[pd.DataFrame] = None

    if not os.path.exists(input_source):
        raise FileNotFoundError(f"Source input file not found at: {input_source}")

    # Case A: Two separate Excel files passed or exist
    if cashflow_source and os.path.exists(cashflow_source):
        input_excel = pd.ExcelFile(input_source)
        contract_sheet = resolve_sheet_name(input_excel.sheet_names, 'Input_Table', 'Contract_Input', 'Contracts', 'Input', 'Master') or input_excel.sheet_names[0]
        input_df = input_excel.parse(contract_sheet)

        cashflow_excel = pd.ExcelFile(cashflow_source)
        cashflow_sheet = resolve_sheet_name(cashflow_excel.sheet_names, 'Cashflow_Table', 'Cashflow_Contracts', 'Cashflow', 'Output') or cashflow_excel.sheet_names[0]
        cashflow_df = cashflow_excel.parse(cashflow_sheet)

    # Case B: Single Excel workbook containing both sheets
    else:
        workbook = pd.ExcelFile(input_source)
        sheet_names = workbook.sheet_names

        contract_sheet = resolve_sheet_name(sheet_names, 'Input_Table', 'Contract_Input', 'Contracts', 'Input')
        cashflow_sheet = resolve_sheet_name(sheet_names, 'Cashflow_Table', 'Cashflow_Contracts', 'Cashflow', 'Output')

        if contract_sheet and cashflow_sheet:
            input_df = workbook.parse(contract_sheet)
            cashflow_df = workbook.parse(cashflow_sheet)
        else:
            raise ValueError(
                f"Could not locate both contract and cashflow data. Provided file: {input_source}, cashflow file: {cashflow_source}"
            )

    # Standardize column header mappings
    contract_id_col = find_column(input_df, ['CONTRACT_ID', 'Contract ID', 'Contract', 'ID'])
    start_date_col = find_column(input_df, ['INTEREST_START_DATE', 'START_DATE', 'Start Date', 'Effective Date', 'Contract Date'])
    principal_col = find_column(input_df, ['CURRENT_PRINCIPAL', 'PRINCIPAL', 'Principal', 'Original Principal', 'Amount', 'Balance'])
    rate_col = find_column(input_df, ['FEE_RATE', 'INTEREST_RATE', 'Interest Rate', 'Rate', 'Annual Rate', 'Coupon'])
    freq_col = find_column(input_df, ['INTEREST_FREQUENCY', 'Interest Frequency', 'Calculation Frequency', 'Frequency', 'DAY_COUNT', 'Day Count'])

    if not contract_id_col or not principal_col or not rate_col:
        raise ValueError(
            f"Input master data missing required columns. Found: {list(input_df.columns)}. "
            "Required: Contract ID, Principal, Interest Rate."
        )

    # Rename master df columns to canonical names without causing duplicate columns
    mapping_tuples = [
        (contract_id_col, 'CONTRACT_ID'),
        (start_date_col, 'START_DATE'),
        (principal_col, 'PRINCIPAL'),
        (rate_col, 'INTEREST_RATE'),
        (freq_col, 'INTEREST_FREQUENCY')
    ]
    for src, tgt in mapping_tuples:
        if src and src != tgt:
            if tgt in input_df.columns:
                input_df = input_df.drop(columns=[tgt])
            input_df = input_df.rename(columns={src: tgt})

    input_df = input_df.loc[:, ~input_df.columns.duplicated()]

    # Normalize cashflow df columns
    cf_contract_id_col = find_column(cashflow_df, ['UNIQUE_ID', 'CONTRACT_ID', 'Contract ID', 'Contract', 'ID'])
    cf_date_col = find_column(cashflow_df, ['FLOW_PAYMENT_DATE', 'PAYMENT_DATE', 'Payment Date', 'Date', 'Cashflow Date', 'Inflow Date'])
    cf_reported_interest_col = find_column(
        cashflow_df,
        ['AMOUNT', 'REPORTED_INTEREST', 'Vendor Interest Amount', 'Reported Interest', 'Recorded Interest', 'Interest Amount', 'Interest']
    )

    if not cf_contract_id_col or not cf_date_col:
        raise ValueError(
            f"Cashflow data missing required columns. Found: {list(cashflow_df.columns)}. "
            "Required: Contract ID and Payment Date."
        )

    cf_mapping_tuples = [
        (cf_contract_id_col, 'CONTRACT_ID'),
        (cf_date_col, 'PAYMENT_DATE'),
        (cf_reported_interest_col, 'REPORTED_INTEREST')
    ]
    for src, tgt in cf_mapping_tuples:
        if src and src != tgt:
            if tgt in cashflow_df.columns:
                cashflow_df = cashflow_df.drop(columns=[tgt])
            cashflow_df = cashflow_df.rename(columns={src: tgt})

    cashflow_df = cashflow_df.loc[:, ~cashflow_df.columns.duplicated()]

    if 'REPORTED_INTEREST' not in cashflow_df.columns:
        cashflow_df['REPORTED_INTEREST'] = 0.0

    # Clean and normalize contract IDs as strings (handling float representations like 5390792549.0)
    input_df['CONTRACT_ID'] = input_df['CONTRACT_ID'].apply(normalize_contract_id)
    cashflow_df['CONTRACT_ID'] = cashflow_df['CONTRACT_ID'].apply(normalize_contract_id)

    # Filter for specific target_contract_id if requested
    if target_contract_id and str(target_contract_id).strip().upper() != 'ALL':
        tid = normalize_contract_id(target_contract_id)
        input_df = input_df[input_df['CONTRACT_ID'] == tid].copy()
        cashflow_df = cashflow_df[cashflow_df['CONTRACT_ID'] == tid].copy()

    return input_df, cashflow_df


def load_contract_data_from_file(file_path: str, contract_id: str):
    """Backward-compatible helper wrapping load_financial_data."""
    return load_financial_data(input_source=file_path, cashflow_source=None, target_contract_id=contract_id)


# ---------------------------------------------------------
# 3. Financial Computation Engine (Flat Interest Calculation)
# ---------------------------------------------------------
def calculate_schedule_for_contract(
    contract_row: pd.Series,
    cashflow_rows: pd.DataFrame
) -> pd.DataFrame:
    """
    Computes flat interest accrual and schedule for a single contract matching exact user specs:
    Columns: CONTRACT_ID, START_DATE, END_DATE, DAYS, OUTSTANDING_PRINCIPAL, INTEREST_RATE, INTEREST_CALCULATED, ACTUAL_INTEREST, DIFFERENCE
    """
    def _scalar(val):
        if isinstance(val, pd.Series):
            return val.iloc[0]
        return val

    contract_id = _scalar(contract_row['CONTRACT_ID'])
    principal = float(_scalar(contract_row['PRINCIPAL']))
    raw_rate = float(_scalar(contract_row['INTEREST_RATE']))

    # Rate handling: convert percentage to decimal format (e.g., 1.0 or 5.5 -> 0.01 or 0.055)
    if raw_rate >= 1.0:
        rate_decimal = raw_rate / 100.0
    else:
        rate_decimal = raw_rate

    freq_val = contract_row.get('INTEREST_FREQUENCY', None)
    freq_days = get_frequency_days(freq_val)

    # Extract only matching cashflow rows for this contract ID and sort chronologically
    cf = cashflow_rows[cashflow_rows['CONTRACT_ID'] == contract_id].copy()
    cf['PAYMENT_DATE'] = pd.to_datetime(cf['PAYMENT_DATE'])
    cf = cf.sort_values(by='PAYMENT_DATE').reset_index(drop=True)

    schedule_records = []
    previous_date: Optional[pd.Timestamp] = None

    for idx, row in cf.iterrows():
        end_date = pd.to_datetime(row['PAYMENT_DATE'])

        if idx == 0:
            # First row: START_DATE = First cashflow date - INTEREST_FREQUENCY
            start_date = end_date - pd.DateOffset(days=freq_days)
        else:
            # Subsequent rows: START_DATE = Previous row's END_DATE
            start_date = previous_date

        actual_days = (end_date - start_date).days

        # Smart days logic:
        # If actual_days is close to payment frequency (within 10 days tolerance), stick to freq_days.
        # Otherwise (for stub / short periods like 9d/10d), stick to actual_days.
        if abs(actual_days - freq_days) <= 10:
            days = freq_days
        else:
            days = actual_days

        # Interest calculation formula: P * rate_decimal * (DAYS / 360)
        calculated_interest = round(principal * rate_decimal * (days / 360.0), 2)
        actual_interest = float(row.get('REPORTED_INTEREST', 0.0) or 0.0)
        difference = round(calculated_interest - actual_interest, 2)

        record = {
            'CONTRACT_ID': contract_id,
            'START_DATE': start_date.strftime('%Y-%m-%d'),
            'END_DATE': end_date.strftime('%Y-%m-%d'),
            'DAYS': days,
            'OUTSTANDING_PRINCIPAL': principal,
            'INTEREST_RATE': rate_decimal,
            'INTEREST_CALCULATED': calculated_interest,
            'ACTUAL_INTEREST': actual_interest,
            'DIFFERENCE': difference
        }
        schedule_records.append(record)
        previous_date = end_date

    return pd.DataFrame(schedule_records)


def calculate_schedule(contract_df: pd.DataFrame, cashflow_df: pd.DataFrame) -> pd.DataFrame:
    """Calculates interest schedule across all contracts in contract_df."""
    if contract_df.empty:
        raise ValueError("Contract dataset is empty.")
    if cashflow_df.empty:
        raise ValueError("No cash flow records found.")

    all_schedules = []
    for _, contract_row in contract_df.iterrows():
        cid = contract_row['CONTRACT_ID']
        c_cashflows = cashflow_df[cashflow_df['CONTRACT_ID'] == cid]
        if not c_cashflows.empty:
            sched = calculate_schedule_for_contract(contract_row, c_cashflows)
            all_schedules.append(sched)

    if not all_schedules:
        raise ValueError("No matching cash flow events found for the provided Contract ID(s).")

    return pd.concat(all_schedules, ignore_index=True)


# ---------------------------------------------------------
# 4. Formatted Excel Generation Engine (openpyxl)
# ---------------------------------------------------------
def export_to_excel(
    contract_df: pd.DataFrame,
    schedule_df: pd.DataFrame,
    output_filename: Optional[str] = None
) -> str:
    """Generates a professionally styled Excel workbook matching exact required columns."""
    if output_filename is None:
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        if len(contract_df) == 1:
            cid = contract_df.iloc[0]['CONTRACT_ID']
            output_filename = f"Contract_Schedule_{cid}_{timestamp}.xlsx"
        else:
            output_filename = f"Financial_Reconciliation_Report_{timestamp}.xlsx"

    with pd.ExcelWriter(output_filename, engine='openpyxl') as writer:
        schedule_df.to_excel(writer, sheet_name='Interest_Schedule', index=False)

        workbook = writer.book
        ws_sched = workbook['Interest_Schedule']

        # Styling
        navy_header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

        alert_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
        alert_font = Font(name="Calibri", size=11, color="C00000", bold=True)

        total_row_font = Font(name="Calibri", size=11, bold=True)
        double_bottom = Side(border_style="double", color="000000")
        top_thin = Side(border_style="thin", color="000000")
        total_border = Border(top=top_thin, bottom=double_bottom)

        for col_idx, cell in enumerate(ws_sched[1], 1):
            cell.fill = navy_header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        num_rows = len(schedule_df)
        for row_idx in range(2, num_rows + 2):
            # Alignment & Number formats
            ws_sched.cell(row=row_idx, column=1).alignment = Alignment(horizontal="center")
            ws_sched.cell(row=row_idx, column=2).alignment = Alignment(horizontal="center")
            ws_sched.cell(row=row_idx, column=3).alignment = Alignment(horizontal="center")

            # DAYS
            ws_sched.cell(row=row_idx, column=4).number_format = '#,##0'
            ws_sched.cell(row=row_idx, column=4).alignment = Alignment(horizontal="right")

            # OUTSTANDING_PRINCIPAL
            ws_sched.cell(row=row_idx, column=5).number_format = '$#,##0.00'

            # INTEREST_RATE - Decimal format (0.01 displays as 1.00%)
            ws_sched.cell(row=row_idx, column=6).number_format = '0.00%'

            # INTEREST_CALCULATED, ACTUAL_INTEREST & DIFFERENCE
            ws_sched.cell(row=row_idx, column=7).number_format = '$#,##0.00'
            ws_sched.cell(row=row_idx, column=8).number_format = '$#,##0.00'
            ws_sched.cell(row=row_idx, column=9).number_format = '$#,##0.00'

            # Alert formatting if difference != 0
            diff_val = ws_sched.cell(row=row_idx, column=9).value
            if diff_val is not None and abs(float(diff_val)) >= 0.01:
                ws_sched.cell(row=row_idx, column=9).fill = alert_fill
                ws_sched.cell(row=row_idx, column=9).font = alert_font

        # Total Summary Row
        tot_row_idx = num_rows + 2
        ws_sched.cell(row=tot_row_idx, column=1, value="Total").font = total_row_font
        ws_sched.cell(row=tot_row_idx, column=7, value=f"=SUM(G2:G{tot_row_idx-1})").number_format = '$#,##0.00'
        ws_sched.cell(row=tot_row_idx, column=8, value=f"=SUM(H2:H{tot_row_idx-1})").number_format = '$#,##0.00'
        ws_sched.cell(row=tot_row_idx, column=9, value=f"=SUM(I2:I{tot_row_idx-1})").number_format = '$#,##0.00'

        for col_idx in range(1, 10):
            c = ws_sched.cell(row=tot_row_idx, column=col_idx)
            c.font = total_row_font
            c.border = total_border

        # Auto-fit Column Widths
        for col in ws_sched.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws_sched.column_dimensions[col_letter].width = max(max_len + 5, 16)

    print(f"[+] Output generated successfully: {output_filename}")
    return output_filename


# ---------------------------------------------------------
# 5. Orchestrator & Entry Point
# ---------------------------------------------------------
def process_contract_file(
    input_file: str = "BSTRAT_Cashflow_Analysis.xlsx",
    cashflow_file: str = "Output_BSTRAT_Cashflow.xlsx",
    contract_id: Optional[str] = None
) -> str:
    """Processes contracts and cashflows from input workbooks and generates output Excel schedule."""
    print(f"[*] Loading master input file: {input_file}")
    if os.path.exists(cashflow_file):
        print(f"[*] Loading cashflow events file: {cashflow_file}")

    if contract_id and str(contract_id).strip().upper() != 'ALL':
        print(f"[*] Filtering for Contract ID: {contract_id}")
    else:
        print("[*] Processing all contracts found in data sources.")

    contract_df, cashflow_df = load_financial_data(
        input_source=input_file,
        cashflow_source=cashflow_file,
        target_contract_id=contract_id
    )

    processed_schedule = calculate_schedule(contract_df, cashflow_df)
    output_filename = export_to_excel(contract_df, processed_schedule)

    return output_filename


if __name__ == "__main__":
    print("=" * 65)
    print(" FINANCIAL CASHFLOW RECONCILIATION & INTEREST CALCULATOR ")
    print("=" * 65)

    default_input = "BSTRAT_Cashflow_Analysis.xlsx"
    default_cashflow = "Output_BSTRAT_Cashflow.xlsx"

    try:
        process_contract_file(
            input_file="BSTRAT_Cashflow_Analysis.xlsx",
            cashflow_file="Output_BSTRAT_Cashflow.xlsx",
            contract_id="5390792549"
        )
    except Exception as e:
        print(f"[!] Error processing financial files: {e}")
        import traceback
        traceback.print_exc()
